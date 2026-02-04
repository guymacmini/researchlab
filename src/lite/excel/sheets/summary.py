"""Summary sheet for institutional research workbook."""

from typing import Dict, Any, List
from datetime import datetime
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


class SummarySheet:
    """Creates executive summary sheet for research report."""
    
    def __init__(self, workbook: Workbook):
        self.workbook = workbook
        self.ws = workbook.create_sheet("Summary", 0)
    
    def create_sheet(self, query: str, analysis: Dict[str, Any], companies: List[Dict[str, str]]):
        """Create the summary sheet with key findings and recommendations."""
        
        # Sheet title
        self.ws['A1'] = "EXECUTIVE SUMMARY"
        self.ws['A1'].style = 'header_style'
        self.ws.merge_cells('A1:H1')
        
        # Research query
        self.ws['A3'] = "Research Query:"
        self.ws['A3'].font = Font(bold=True, size=12)
        self.ws['A4'] = query
        self.ws['A4'].font = Font(size=11)
        self.ws.merge_cells('A4:H4')
        
        # Timestamp
        self.ws['A6'] = f"Report Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        self.ws['A6'].font = Font(size=10, italic=True)
        
        # Companies analyzed
        self.ws['A8'] = "COMPANIES ANALYZED"
        self.ws['A8'].style = 'subheader_style'
        self.ws.merge_cells('A8:D8')
        
        row = 9
        # Headers for companies table
        headers = ['Ticker', 'Company Name', 'Investment Rationale']
        for col, header in enumerate(headers, 1):
            cell = self.ws.cell(row=row, column=col, value=header)
            cell.style = 'subheader_style'
        row += 1
        
        # Company data
        for company in companies:
            self.ws.cell(row=row, column=1, value=company.get('ticker', 'N/A'))
            self.ws.cell(row=row, column=2, value=company.get('name', 'Unknown'))
            self.ws.cell(row=row, column=3, value=company.get('rationale', 'No rationale provided'))
            
            # Style the row
            for col in range(1, 4):
                self.ws.cell(row=row, column=col).style = 'data_style'
            row += 1
        
        # Extract key findings from analysis text
        analysis_text = analysis.get('analysis_text', '')
        key_findings = self._extract_key_findings(analysis_text)
        
        # Key findings section
        row += 2
        self.ws[f'A{row}'] = "KEY FINDINGS & RECOMMENDATIONS"
        self.ws[f'A{row}'].style = 'subheader_style'
        self.ws.merge_cells(f'A{row}:H{row}')
        row += 1
        
        for finding in key_findings:
            self.ws[f'A{row}'] = f"• {finding}"
            self.ws[f'A{row}'].font = Font(size=11)
            self.ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
            self.ws.merge_cells(f'A{row}:H{row}')
            self.ws.row_dimensions[row].height = 40  # Increase row height for wrapped text
            row += 1
        
        # Investment summary table
        row += 2
        self.ws[f'A{row}'] = "INVESTMENT SUMMARY"
        self.ws[f'A{row}'].style = 'subheader_style'
        self.ws.merge_cells(f'A{row}:F{row}')
        row += 1
        
        # Summary table headers
        summary_headers = ['Ticker', 'Rating', 'Target Price', 'Upside/Downside', 'Risk Level', 'Catalyst']
        for col, header in enumerate(summary_headers, 1):
            cell = self.ws.cell(row=row, column=col, value=header)
            cell.style = 'subheader_style'
        row += 1
        
        # Extract investment recommendations
        recommendations = self._extract_recommendations(analysis_text, companies)
        for rec in recommendations:
            for col, value in enumerate(rec, 1):
                cell = self.ws.cell(row=row, column=col, value=value)
                cell.style = 'data_style'
            row += 1
        
        # Risk factors
        row += 2
        self.ws[f'A{row}'] = "KEY RISK FACTORS"
        self.ws[f'A{row}'].style = 'subheader_style'
        self.ws.merge_cells(f'A{row}:H{row}')
        row += 1
        
        risk_factors = self._extract_risk_factors(analysis_text)
        for risk in risk_factors:
            self.ws[f'A{row}'] = f"• {risk}"
            self.ws[f'A{row}'].font = Font(size=11)
            self.ws[f'A{row}'].alignment = Alignment(wrap_text=True, vertical='top')
            self.ws.merge_cells(f'A{row}:H{row}')
            self.ws.row_dimensions[row].height = 30
            row += 1
        
        # Data sources note
        row += 2
        self.ws[f'A{row}'] = "Data Sources: Financial Modeling Prep, Alpha Vantage, SEC EDGAR, Finnhub"
        self.ws[f'A{row}'].font = Font(size=9, italic=True, color='666666')
        
        # Auto-adjust column widths
        self._adjust_column_widths()
    
    def _extract_key_findings(self, analysis_text: str) -> List[str]:
        """Extract key findings from analysis text."""
        findings = []
        
        # Look for thesis or summary sections
        thesis_patterns = [
            r'(?i)thesis[:\s]+([^\.]+\.)',
            r'(?i)key findings[:\s]+([^\.]+\.)',
            r'(?i)investment thesis[:\s]+([^\.]+\.)',
            r'(?i)summary[:\s]+([^\.]+\.)'
        ]
        
        for pattern in thesis_patterns:
            matches = re.findall(pattern, analysis_text)
            findings.extend(matches[:2])  # Limit to 2 per pattern
        
        # Look for bullet points or numbered items
        bullet_patterns = [
            r'(?i)[-•*]\s*([^\.]+\.)',
            r'(?i)\d+\.\s*([^\.]+\.)'
        ]
        
        for pattern in bullet_patterns:
            matches = re.findall(pattern, analysis_text)
            findings.extend(matches[:3])
        
        # Look for recommendation statements
        recommendation_patterns = [
            r'(?i)recommend[s]?\s+([^\.]+\.)',
            r'(?i)rating[:\s]+([^\.]+\.)',
            r'(?i)target[:\s]+([^\.]+\.)'
        ]
        
        for pattern in recommendation_patterns:
            matches = re.findall(pattern, analysis_text)
            findings.extend(matches[:2])
        
        # Clean and deduplicate findings
        cleaned_findings = []
        for finding in findings:
            finding = finding.strip()
            if len(finding) > 20 and finding not in cleaned_findings:
                cleaned_findings.append(finding)
        
        # If no findings found, provide generic ones
        if not cleaned_findings:
            cleaned_findings = [
                "Investment analysis completed for specified companies",
                "Comprehensive financial and valuation analysis performed",
                "Risk factors and scenario analysis included"
            ]
        
        return cleaned_findings[:5]  # Limit to 5 key findings
    
    def _extract_recommendations(self, analysis_text: str, companies: List[Dict[str, str]]) -> List[List[str]]:
        """Extract investment recommendations for each company."""
        recommendations = []
        
        for company in companies:
            ticker = company.get('ticker', '')
            
            # Try to find company-specific recommendations
            rating = self._extract_rating(analysis_text, ticker)
            target_price = self._extract_target_price(analysis_text, ticker)
            upside = self._calculate_upside(target_price)
            risk_level = self._extract_risk_level(analysis_text, ticker)
            catalyst = self._extract_catalyst(analysis_text, ticker)
            
            recommendations.append([
                ticker,
                rating,
                target_price,
                upside,
                risk_level,
                catalyst
            ])
        
        return recommendations
    
    def _extract_rating(self, text: str, ticker: str) -> str:
        """Extract investment rating for a ticker."""
        # Look for ratings near the ticker
        patterns = [
            rf'(?i){ticker}[^\.]*?(buy|sell|hold)',
            r'(?i)rating[:\s]+(buy|sell|hold)',
            r'(?i)recommend[a-z]*[:\s]+(buy|sell|hold)'
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return match.group(1).upper()
        
        return 'HOLD'  # Default rating
    
    def _extract_target_price(self, text: str, ticker: str) -> str:
        """Extract target price for a ticker."""
        patterns = [
            rf'(?i){ticker}[^\.]*?target[:\s]+\$?(\d+(?:\.\d+)?)',
            r'(?i)target[:\s]+\$?(\d+(?:\.\d+)?)',
            r'(?i)price target[:\s]+\$?(\d+(?:\.\d+)?)'
        ]
        
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                return f"${match.group(1)}"
        
        return "N/A"
    
    def _calculate_upside(self, target_price: str) -> str:
        """Calculate upside/downside from target price."""
        if target_price == "N/A":
            return "N/A"
        
        # This would need current price data to calculate
        # For now, return placeholder
        return "TBD"
    
    def _extract_risk_level(self, text: str, ticker: str) -> str:
        """Extract risk level assessment."""
        if 'high risk' in text.lower():
            return 'HIGH'
        elif 'low risk' in text.lower():
            return 'LOW'
        elif 'moderate' in text.lower():
            return 'MODERATE'
        else:
            return 'MODERATE'
    
    def _extract_catalyst(self, text: str, ticker: str) -> str:
        """Extract key catalyst for the investment."""
        catalyst_patterns = [
            r'(?i)catalyst[:\s]+([^\.]+\.)',
            r'(?i)next[^\.]*?(earnings|quarter|year)',
            r'(?i)expect[^\.]*?(growth|expansion|launch)'
        ]
        
        for pattern in catalyst_patterns:
            match = re.search(pattern, text)
            if match:
                return match.group(1)[:50] + "..." if len(match.group(1)) > 50 else match.group(1)
        
        return "Quarterly earnings"
    
    def _extract_risk_factors(self, analysis_text: str) -> List[str]:
        """Extract key risk factors from analysis."""
        risks = []
        
        # Look for risk sections
        risk_patterns = [
            r'(?i)risk[s]?[:\s]+([^\.]+\.)',
            r'(?i)concern[s]?[:\s]+([^\.]+\.)',
            r'(?i)threat[s]?[:\s]+([^\.]+\.)',
            r'(?i)challenge[s]?[:\s]+([^\.]+\.)'
        ]
        
        for pattern in risk_patterns:
            matches = re.findall(pattern, analysis_text)
            risks.extend(matches[:2])
        
        # Clean risks
        cleaned_risks = []
        for risk in risks:
            risk = risk.strip()
            if len(risk) > 15 and risk not in cleaned_risks:
                cleaned_risks.append(risk)
        
        # Default risks if none found
        if not cleaned_risks:
            cleaned_risks = [
                "Market volatility and macroeconomic conditions",
                "Company-specific operational risks",
                "Sector-specific regulatory or competitive risks"
            ]
        
        return cleaned_risks[:4]  # Limit to 4 risks
    
    def _adjust_column_widths(self):
        """Auto-adjust column widths for better readability."""
        for column in self.ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set width with some padding, max 60
            adjusted_width = min(max_length + 5, 60)
            self.ws.column_dimensions[column_letter].width = max(adjusted_width, 12)