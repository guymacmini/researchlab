"""Excel workbook generator for institutional research reports."""

import io
from typing import Dict, Any, List, Optional, Tuple
from datetime import datetime, date
import structlog

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, NamedStyle
from openpyxl.styles.numbers import FORMAT_CURRENCY_USD, FORMAT_PERCENTAGE, FORMAT_NUMBER_COMMA_SEPARATED1
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

from .sheets.summary import SummarySheet
from .sheets.financial_statements import FinancialStatementsSheet
from .sheets.valuation_comps import ValuationCompsSheet  
from .sheets.scenario_analysis import ScenarioAnalysisSheet
from .sheets.company_profiles import CompanyProfilesSheet
from .sheets.data_sources import DataSourcesSheet

logger = structlog.get_logger()


class InstitutionalWorkbook:
    """
    Professional Excel workbook generator for institutional research.
    
    Creates a comprehensive workbook with:
    - Executive Summary
    - Financial Statements (3-5 years)
    - Valuation Comparables
    - Scenario Analysis (Bull/Base/Bear)
    - Company Profiles
    - Data Sources & Methodology
    """
    
    def __init__(self):
        self.workbook = Workbook()
        self.setup_styles()
        
        # Remove default worksheet
        if 'Sheet' in [ws.title for ws in self.workbook.worksheets]:
            self.workbook.remove(self.workbook['Sheet'])
    
    def setup_styles(self):
        """Setup consistent styling for the workbook."""
        
        # Header styles
        self.header_style = NamedStyle(name="header_style")
        self.header_style.font = Font(name='Calibri', size=14, bold=True, color="FFFFFF")
        self.header_style.fill = PatternFill(start_color="2F4F4F", end_color="2F4F4F", fill_type="solid")
        self.header_style.alignment = Alignment(horizontal="center", vertical="center")
        self.header_style.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Subheader style
        self.subheader_style = NamedStyle(name="subheader_style")
        self.subheader_style.font = Font(name='Calibri', size=12, bold=True)
        self.subheader_style.fill = PatternFill(start_color="E6E6FA", end_color="E6E6FA", fill_type="solid")
        self.subheader_style.alignment = Alignment(horizontal="center", vertical="center")
        self.subheader_style.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Data style
        self.data_style = NamedStyle(name="data_style")
        self.data_style.font = Font(name='Calibri', size=11)
        self.data_style.alignment = Alignment(horizontal="right", vertical="center")
        self.data_style.border = Border(
            left=Side(style='hair'), right=Side(style='hair'),
            top=Side(style='hair'), bottom=Side(style='hair')
        )
        
        # Currency style
        self.currency_style = NamedStyle(name="currency_style")
        self.currency_style.font = Font(name='Calibri', size=11)
        self.currency_style.alignment = Alignment(horizontal="right", vertical="center")
        self.currency_style.number_format = FORMAT_CURRENCY_USD
        self.currency_style.border = Border(
            left=Side(style='hair'), right=Side(style='hair'),
            top=Side(style='hair'), bottom=Side(style='hair')
        )
        
        # Percentage style
        self.percentage_style = NamedStyle(name="percentage_style")
        self.percentage_style.font = Font(name='Calibri', size=11)
        self.percentage_style.alignment = Alignment(horizontal="right", vertical="center")
        self.percentage_style.number_format = FORMAT_PERCENTAGE
        self.percentage_style.border = Border(
            left=Side(style='hair'), right=Side(style='hair'),
            top=Side(style='hair'), bottom=Side(style='hair')
        )
        
        # Add styles to workbook
        self.workbook.add_named_style(self.header_style)
        self.workbook.add_named_style(self.subheader_style)
        self.workbook.add_named_style(self.data_style)
        self.workbook.add_named_style(self.currency_style)
        self.workbook.add_named_style(self.percentage_style)
    
    def create_research_report(self, research_data: Dict[str, Any]) -> io.BytesIO:
        """
        Create complete institutional research report.
        
        Args:
            research_data: Dictionary containing all research analysis results
            
        Returns:
            BytesIO object containing the Excel workbook
        """
        logger.info("Creating institutional research workbook")
        
        try:
            # Extract key components from research data
            query = research_data.get('original_query', 'Investment Research')
            companies = research_data.get('companies_analyzed', [])
            analysis = research_data.get('analysis', {})
            company_data = analysis.get('company_data', {})
            
            # Create sheets in order
            sheets_created = []
            
            # 1. Summary Sheet
            try:
                summary_sheet = SummarySheet(self.workbook)
                summary_sheet.create_sheet(query, analysis, companies)
                sheets_created.append('Summary')
            except Exception as e:
                logger.error("Failed to create summary sheet", error=str(e))
            
            # 2. Financial Statements
            try:
                financial_sheet = FinancialStatementsSheet(self.workbook)
                financial_sheet.create_sheet(company_data)
                sheets_created.append('Financial Statements')
            except Exception as e:
                logger.error("Failed to create financial statements sheet", error=str(e))
            
            # 3. Valuation Comparables
            try:
                valuation_sheet = ValuationCompsSheet(self.workbook)
                valuation_sheet.create_sheet(company_data, companies)
                sheets_created.append('Valuation Comps')
            except Exception as e:
                logger.error("Failed to create valuation comps sheet", error=str(e))
            
            # 4. Scenario Analysis
            try:
                scenario_sheet = ScenarioAnalysisSheet(self.workbook)
                scenario_sheet.create_sheet(company_data, analysis.get('analysis_text', ''))
                sheets_created.append('Scenario Analysis')
            except Exception as e:
                logger.error("Failed to create scenario analysis sheet", error=str(e))
            
            # 5. Company Profiles
            try:
                profiles_sheet = CompanyProfilesSheet(self.workbook)
                profiles_sheet.create_sheet(company_data, companies)
                sheets_created.append('Company Profiles')
            except Exception as e:
                logger.error("Failed to create company profiles sheet", error=str(e))
            
            # 6. Data Sources
            try:
                sources_sheet = DataSourcesSheet(self.workbook)
                sources_sheet.create_sheet(analysis.get('data_sources', []), research_data.get('timestamp'))
                sheets_created.append('Data Sources')
            except Exception as e:
                logger.error("Failed to create data sources sheet", error=str(e))
            
            # Set the summary sheet as active
            if 'Summary' in [ws.title for ws in self.workbook.worksheets]:
                self.workbook.active = self.workbook['Summary']
            
            # Save to BytesIO
            excel_buffer = io.BytesIO()
            self.workbook.save(excel_buffer)
            excel_buffer.seek(0)
            
            logger.info("Research workbook created successfully", 
                       sheets_created=sheets_created,
                       companies_count=len(companies))
            
            return excel_buffer
            
        except Exception as e:
            logger.error("Failed to create research workbook", error=str(e))
            raise
    
    def create_company_analysis(self, company_data: Dict[str, Any]) -> io.BytesIO:
        """Create single-company analysis workbook."""
        logger.info("Creating single company analysis workbook")
        
        try:
            # Create simplified workbook for single company
            ws = self.workbook.create_sheet("Company Analysis", 0)
            
            # Add company header
            symbol = company_data.get('symbol', 'N/A')
            company_name = company_data.get('profile', {}).get('name', 'Unknown Company')
            
            ws['A1'] = f"{company_name} ({symbol}) - Investment Analysis"
            ws['A1'].style = self.header_style
            ws.merge_cells('A1:H1')
            
            # Add timestamp
            ws['A2'] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
            ws['A2'].font = Font(size=10, italic=True)
            
            # Add basic metrics table
            if 'metrics' in company_data:
                metrics = company_data['metrics'].get('metric', {})
                
                row = 4
                ws[f'A{row}'] = "Key Metrics"
                ws[f'A{row}'].style = self.subheader_style
                ws.merge_cells(f'A{row}:B{row}')
                row += 1
                
                key_metrics = [
                    ('P/E Ratio', metrics.get('peBasicExclExtraTTM', 'N/A')),
                    ('P/B Ratio', metrics.get('pbQuarterly', 'N/A')),
                    ('ROE %', metrics.get('roeTTM', 'N/A')),
                    ('Revenue Growth %', metrics.get('revenueGrowthTTMYoy', 'N/A')),
                    ('Gross Margin %', metrics.get('grossMarginTTM', 'N/A')),
                ]
                
                for metric_name, value in key_metrics:
                    ws[f'A{row}'] = metric_name
                    ws[f'B{row}'] = value
                    ws[f'B{row}'].style = self.data_style
                    row += 1
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Save to BytesIO
            excel_buffer = io.BytesIO()
            self.workbook.save(excel_buffer)
            excel_buffer.seek(0)
            
            logger.info("Single company workbook created", symbol=symbol)
            return excel_buffer
            
        except Exception as e:
            logger.error("Failed to create company analysis workbook", error=str(e))
            raise
    
    @staticmethod
    def format_number(value: Any) -> Any:
        """Format numbers for Excel display."""
        if value is None or value == 'N/A':
            return 'N/A'
        
        try:
            if isinstance(value, str):
                # Try to convert string to number
                if value.replace('.', '').replace('-', '').replace('+', '').isdigit():
                    value = float(value)
                else:
                    return value
            
            if isinstance(value, (int, float)):
                return float(value)
            
            return value
            
        except (ValueError, AttributeError):
            return value
    
    @staticmethod
    def safe_get(data: Dict[str, Any], key: str, default: Any = 'N/A') -> Any:
        """Safely get value from dictionary with default."""
        try:
            return data.get(key, default)
        except (AttributeError, KeyError):
            return default


def create_institutional_report(research_data: Dict[str, Any]) -> io.BytesIO:
    """
    Convenience function to create institutional research report.
    
    Args:
        research_data: Complete research analysis data
        
    Returns:
        BytesIO containing Excel workbook
    """
    workbook_generator = InstitutionalWorkbook()
    return workbook_generator.create_research_report(research_data)


def create_company_report(company_data: Dict[str, Any]) -> io.BytesIO:
    """
    Convenience function to create single company report.
    
    Args:
        company_data: Single company financial data
        
    Returns:
        BytesIO containing Excel workbook
    """
    workbook_generator = InstitutionalWorkbook()
    return workbook_generator.create_company_analysis(company_data)